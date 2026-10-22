import io
from datetime import UTC, datetime, timedelta

import pymupdf
from docx import Document as WordDocument
from sqlalchemy import select

from app.core.config import get_settings
from app.db.models import IngestionJob, KnowledgeBase, KnowledgeBaseMember, User
from app.worker.pipeline import process_ingestion_job


def kb(client):
    return client.get("/api/knowledge-bases").json()[0]["id"]


def test_crud_scope_cycles_batch_and_trash(client, factory):
    base = kb(client)
    other = client.post("/api/knowledge-bases", json={"name": "另一个库"}).json()["id"]
    folder = client.post(f"/api/knowledge-bases/{base}/folders", json={"name": "论文"}).json()
    child = client.post(f"/api/knowledge-bases/{base}/folders", json={"name": "机器学习", "parent_id": folder["id"]}).json()
    assert client.patch(f"/api/folders/{folder['id']}", json={"parent_id": child["id"]}).status_code == 400
    assert client.post(f"/api/knowledge-bases/{other}/folders", json={"name": "非法", "parent_id": folder["id"]}).status_code == 400
    tag = client.post(f"/api/knowledge-bases/{base}/tags", json={"name": "ＬＬＭ", "aliases": ["大语言模型"]}).json()
    assert client.post(f"/api/knowledge-bases/{base}/tags", json={"name": "llm"}).status_code == 409
    foreign_tag = client.post(f"/api/knowledge-bases/{other}/tags", json={"name": "独立标签"}).json()
    result = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", ("论文.md", "# 检索\n\n向量检索与中文知识管理。".encode(), "text/markdown"))]).json()
    doc = result["items"][0]["document"]
    assert client.post("/api/documents/batch", json={"kb_id": base, "document_ids": [doc["id"]], "action": "tags", "tag_ids": [foreign_tag["id"]]}).status_code == 400
    updated = client.post("/api/documents/batch", json={"kb_id": base, "document_ids": [doc["id"]], "action": "move", "folder_id": child["id"]})
    assert updated.status_code == 200
    updated = client.patch(f"/api/documents/{doc['id']}", json={"tag_ids": [tag["id"]], "version": doc["version"]})
    assert updated.status_code == 409
    assert client.delete(f"/api/documents/{doc['id']}").status_code == 200
    assert client.get(f"/api/documents/{doc['id']}/download").status_code == 404
    assert client.post(f"/api/documents/{doc['id']}/restore").status_code == 200
    assert client.get(f"/api/documents/{doc['id']}/download").content.startswith(b"#")
    assert client.delete(f"/api/knowledge-bases/{other}").status_code == 200
    assert client.get(f"/api/knowledge-bases/{other}").status_code == 404
    assert client.post(f"/api/knowledge-bases/{other}/restore").status_code == 200


def test_private_permissions_and_host_origin(client, factory):
    with factory() as db:
        user = User(email="other@localhost", display_name="Other")
        db.add(user)
        db.flush()
        private = KnowledgeBase(owner_id=user.id, name="私有")
        db.add(private)
        db.flush()
        db.add(KnowledgeBaseMember(kb_id=private.id, user_id=user.id, role="owner"))
        db.commit()
        private_id = private.id
    assert client.get(f"/api/knowledge-bases/{private_id}").status_code == 403
    assert client.post("/api/search", json={"query": "秘密", "kb_ids": [private_id]}).status_code == 403
    assert client.get("/api/health", headers={"Host": "public.example"}).status_code == 403
    assert client.get("/api/health", headers={"Origin": "https://evil.example"}).status_code == 403


def test_real_four_formats_upload_parse_search_and_duplicate(client, factory):
    base = kb(client)
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Knowledge retrieval")
    pdf_data = pdf.tobytes()
    pdf.close()
    word = WordDocument()
    word.add_heading("中文论文", level=1)
    word.add_paragraph("中文检索与知识聚类。")
    buf = io.BytesIO()
    word.save(buf)
    fixtures = [("sample.pdf", pdf_data), ("sample.docx", buf.getvalue()), ("sample.md", "# 中文知识\n\n知识聚类可以提高中文检索效率。".encode()), ("sample.txt", "中文知识聚类\n中文检索测试。".encode())]
    response = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", (name, data)) for name, data in fixtures])
    assert response.status_code == 200, response.text
    result = response.json()
    assert all(item["status"] == "queued" for item in result["items"]), result
    for job in result["jobs"]:
        process_ingestion_job(job["id"], factory, get_settings())
    for item in result["items"]:
        doc = client.get(f"/api/documents/{item['document']['id']}").json()
        assert doc["parse_status"] == "ready", doc
        assert doc["index_status"] == "unconfigured"
        sections = client.get(f"/api/documents/{doc['id']}/sections").json()
        assert sections and sections[0]["anchors"]
        if doc["ext"] == "docx":
            assert all("page_start" not in section["anchors"] for section in sections)
    found = client.post("/api/search", json={"query": "中文检索", "kb_ids": [base]}).json()
    assert found["items"] and found["mode"] == "keyword", found
    duplicate = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", fixtures[2])]).json()
    assert duplicate["items"][0]["status"] == "duplicate"
    copied = client.post(f"/api/knowledge-bases/{base}/uploads", data={"duplicate_policy": "copy"}, files=[("files", fixtures[2])]).json()
    assert copied["items"][0]["status"] == "queued"
    invalid = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", ("bad.pdf", b"not pdf")), ("files", ("../bad.txt", b"path")), ("files", ("code.exe", b"MZ"))]).json()
    assert all(item["status"] == "error" for item in invalid["items"])
    assert "api_key" not in client.get("/api/settings/model-status").text.lower()


def test_upload_size_and_mime_limits(client, monkeypatch):
    base = kb(client)
    monkeypatch.setattr(get_settings(), "max_upload_size_mb", 1)
    result = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", ("huge.txt", b"x" * (1024*1024+1), "text/plain")), ("files", ("mismatch.md", b"# text", "image/png"))]).json()
    assert all(item["status"] == "error" for item in result["items"])
    assert "限制" in result["items"][0]["message"]
    assert "类型" in result["items"][1]["message"]
    oversized = client.post(f"/api/knowledge-bases/{base}/uploads", content=b"", headers={"Content-Length": str(300*1024*1024)})
    assert oversized.status_code == 413


def test_transaction_commit_failure_cannot_return_success(client, factory):
    from sqlalchemy.exc import IntegrityError

    from app.db.session import get_db
    from app.main import app
    original = app.dependency_overrides[get_db]
    def failing_commit():
        with factory() as db:
            yield db
            db.rollback()
            raise IntegrityError("commit failed", {}, RuntimeError("simulated failure"))
    app.dependency_overrides[get_db] = failing_commit
    try:
        response = client.post("/api/knowledge-bases", json={"name": "不应返回成功"})
        assert response.status_code == 409
        with factory() as db:
            assert db.scalar(select(KnowledgeBase).where(KnowledgeBase.name == "不应返回成功")) is None
    finally:
        app.dependency_overrides[get_db] = original


def test_conversations_persistent_no_key_and_search_isolation(client, factory):
    base = kb(client)
    conversation = client.post("/api/conversations", json={"scope": {"kb_ids": [base]}}).json()
    assert client.post(f"/api/conversations/{conversation['id']}/messages", json={"content": "资料里有没有量子力学？"}).status_code == 202
    from app.services.chat import process_chat_job
    with factory() as db:
        job = db.scalar(select(IngestionJob).where(IngestionJob.kind == "chat"))
        job_id = job.id
    process_chat_job(job_id, factory, get_settings())
    messages = client.get(f"/api/conversations/{conversation['id']}/messages").json()
    assert len(messages) == 2 and "找不到足够依据" in messages[-1]["content"]
    assert client.get(f"/api/conversations/{conversation['id']}").json()["scope"]["kb_ids"] == [base]


def test_conversation_message_order_is_monotonic_with_frozen_clock(client, factory, monkeypatch):
    import app.api.workflows as workflows
    from app.db.models import Message
    from app.services.chat import process_chat_job

    frozen = datetime(2026, 10, 22, 12, tzinfo=UTC)
    monkeypatch.setattr(workflows, "utcnow", lambda: frozen)
    monkeypatch.setattr(Message.__table__.c.created_at.default, "arg", lambda _context: frozen)
    # Reverse IDs ensure a same-timestamp regression fails deterministically.
    ids = iter([f"00000000-0000-0000-0000-{number:012d}" for number in (4, 3, 2, 1)])
    monkeypatch.setattr(Message.__table__.c.id.default, "arg", lambda _context: next(ids))
    base = kb(client)
    created = client.post("/api/conversations", json={"scope": {"kb_ids": [base]}})
    assert created.status_code == 201
    conversation_id = created.json()["id"]
    assistant_ids = []
    for question in ("第一轮问题", "第二轮问题"):
        queued = client.post(f"/api/conversations/{conversation_id}/messages", json={"content": question})
        assert queued.status_code == 202
        assistant_ids.append(queued.json()["id"])
        with factory() as db:
            job_id = next(job.id for job in db.scalars(select(IngestionJob).where(IngestionJob.kind == "chat"))
                          if job.payload["message_id"] == queued.json()["id"])
        process_chat_job(job_id, factory, get_settings())
    messages = client.get(f"/api/conversations/{conversation_id}/messages").json()
    assert [message["role"] for message in messages] == ["user", "assistant", "user", "assistant"]
    assert [message["content"] for message in messages if message["role"] == "user"] == ["第一轮问题", "第二轮问题"]
    assert [message["id"] for message in messages if message["role"] == "assistant"] == assistant_ids
    assert all(message["status"] == "completed" for message in messages)
    timestamps = [datetime.fromisoformat(message["created_at"]).replace(tzinfo=UTC) for message in messages]
    assert timestamps == [frozen + timedelta(microseconds=offset) for offset in range(4)]


def test_rrf_dimension_and_scoped_sql():
    from app.retrieval.search import cosine, rrf
    ranking, scores = rrf([["a", "b", "b"], ["b", "c"]])
    assert ranking[0] == "b" and scores["b"] > scores["a"]
    import pytest
    with pytest.raises(ValueError):
        cosine([1], [1, 2])


def test_chat_citation_allowlist_and_sse_replay(client, factory, monkeypatch):
    from app.ai.gateway import ModelGateway
    from app.db.models import ConversationEvent
    from app.services.chat import process_chat_job
    base = kb(client)
    uploaded = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", ("证据.md", "# 中文检索\n\n中文检索通过关键词和向量融合实现。".encode()))]).json()
    process_ingestion_job(uploaded["jobs"][0]["id"], factory, get_settings())
    conversation = client.post("/api/conversations", json={"scope": {"kb_ids": [base]}}).json()
    queued = client.post(f"/api/conversations/{conversation['id']}/messages", json={"content": "中文检索"}).json()
    monkeypatch.setattr(ModelGateway, "llm_configured", property(lambda _: True))
    monkeypatch.setattr(ModelGateway, "stream_chat", lambda *_: iter(["关键词与向量融合", "[S1]", "[S999]"]))
    with factory() as db:
        job_id = db.scalar(select(IngestionJob.id).where(IngestionJob.kind == "chat"))
    process_chat_job(job_id, factory, get_settings())
    answer = client.get(f"/api/conversations/{conversation['id']}/messages").json()[-1]
    assert answer["id"] == queued["id"] and answer["status"] == "completed"
    assert "[S999]" not in answer["content"] and "[引用无效]" in answer["content"]
    assert len(answer["citations"]) == 1 and answer["citations"][0]["ordinal"] == 1
    assert answer["citations"][0]["document_id"] == uploaded["items"][0]["document"]["id"]
    with factory() as db:
        first = db.scalar(select(ConversationEvent).where(ConversationEvent.conversation_id == conversation["id"]).order_by(ConversationEvent.sequence))
        first_sequence = first.sequence
    response = client.get(f"/api/conversations/{conversation['id']}/events?after={first_sequence}")
    assert response.status_code == 200 and "event: token" in response.text and "event: completed" in response.text
    assert f"id: {first_sequence}\n" not in response.text
    client.delete(f"/api/documents/{answer['citations'][0]['document_id']}")
    stale = client.get(f"/api/conversations/{conversation['id']}/messages").json()[-1]
    assert stale["citations"][0]["available"] is False


def test_chat_lease_transfer_prevents_stale_writes(client, factory, monkeypatch):
    from app.ai.gateway import ModelGateway
    from app.db.models import Message
    from app.services.chat import process_chat_job
    base = kb(client)
    uploaded = client.post(f"/api/knowledge-bases/{base}/uploads", files=[("files", ("资料.txt", "中文检索真实证据".encode()))]).json()
    process_ingestion_job(uploaded["jobs"][0]["id"], factory, get_settings())
    conversation = client.post("/api/conversations", json={"scope": {"kb_ids": [base]}}).json()
    response = client.post(f"/api/conversations/{conversation['id']}/messages", json={"content": "中文检索"}).json()
    with factory() as db:
        job = db.scalar(select(IngestionJob).where(IngestionJob.kind == "chat"))
        job.status, job.lease_owner = "running", "old-worker"
        job_id = job.id
        db.commit()
    def transferred(*_):
        with factory() as db:
            job = db.get(IngestionJob, job_id)
            job.lease_owner = "new-worker"
            db.commit()
        yield "旧worker禁止写入"
    monkeypatch.setattr(ModelGateway, "llm_configured", property(lambda _: True))
    monkeypatch.setattr(ModelGateway, "stream_chat", transferred)
    process_chat_job(job_id, factory, get_settings())
    with factory() as db:
        assert db.get(Message, response["id"]).content == ""
        assert db.get(IngestionJob, job_id).lease_owner == "new-worker"
