"""Real source fixtures for the structured formats, including honest anchors."""

import io
import json
import zipfile
from datetime import date

import pytest

from app.ingestion.formats import (
    parse_csv,
    parse_epub,
    parse_html,
    parse_json,
    parse_jsonl,
    parse_pptx,
    parse_xlsx,
    validate_csv,
    validate_epub,
    validate_json,
    validate_jsonl,
    validate_pptx,
    validate_xlsx,
)
from app.ingestion.formats.common import DocumentBuilder
from app.ingestion.parser_types import ParseError


def presentation_bytes():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "中文幻灯片"
    slide.placeholders[1].text = "真实的演示正文"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "录取名额"
    table.cell(1, 1).text = "28"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)).text_frame.text = "组合图形文字"
    slide.notes_slide.notes_text_frame.text = "讲者备注"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "下一页"
    second.placeholders[1].text = "后续证据"
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def workbook_bytes():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "金额", "日期"])
    sheet.append(["设备", 28, date(2026, 10, 25)])
    sheet.append(["合计", "=SUM(B2:B2)"])
    second = workbook.create_sheet("研发")
    second.append(["阶段", "负责人"])
    second.append(["验证", "学姐"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def epub_bytes(*, encrypted=False, href="chapter-two.xhtml"):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/book.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        archive.writestr("OPS/book.opf", f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><manifest><item id="one" href="chapter-one.xhtml" media-type="application/xhtml+xml"/><item id="two" href="{href}" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="two"/><itemref idref="one"/></spine></package>')
        archive.writestr("OPS/chapter-one.xhtml", '<html><body><h1>第一章</h1><p>先放在包内，但阅读顺序在后。</p></body></html>')
        archive.writestr("OPS/chapter-two.xhtml", '<!DOCTYPE html><html><body><h1>第二章</h1><p>真实阅读顺序在前。</p><script>不应被提取</script><img src="https://example.invalid/tracker.png"/></body></html>')
        if encrypted:
            archive.writestr("META-INF/encryption.xml", "<encryption/>")
    return stream.getvalue()


def assert_source_offsets(parsed):
    for section in parsed.sections:
        assert parsed.plain_text[section.source_start:section.source_end] == section.text


def test_pptx_reads_slides_tables_groups_notes_with_slide_anchors():
    data = presentation_bytes()
    validate_pptx(data)
    parsed = parse_pptx(data)
    assert len(parsed.sections) == 2
    assert parsed.sections[0].anchor == {"slide_start": 1, "slide_end": 1}
    assert parsed.sections[1].anchor == {"slide_start": 2, "slide_end": 2}
    for text in ("中文幻灯片", "真实的演示正文", "录取名额 | 28", "组合图形文字", "讲者备注"):
        assert text in parsed.plain_text
    assert "page_start" not in parsed.sections[0].anchor
    assert_source_offsets(parsed)


def test_xlsx_streams_sheets_dates_formulas_without_fake_pages():
    parsed = parse_xlsx(workbook_bytes())
    assert parsed.metadata["formulas"] == "stored_text_not_evaluated"
    assert "2026-10-25" in parsed.plain_text
    assert "=SUM(B2:B2)" in parsed.plain_text
    assert parsed.sections[0].anchor == {"sheet_index": 1, "sheet_name": "预算", "row_start": 1, "row_end": 3}
    assert parsed.sections[1].anchor["sheet_name"] == "研发"
    assert_source_offsets(parsed)


def test_xlsx_groups_rows_without_losing_real_row_numbers():
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("批次")
    for row in range(1, 207):
        sheet.append([row, f"资料 {row}"])
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = parse_xlsx(stream.getvalue())
    assert [(section.anchor["row_start"], section.anchor["row_end"]) for section in parsed.sections] == [(1, 100), (101, 200), (201, 206)]
    assert "资料 206" in parsed.plain_text
    assert_source_offsets(parsed)


def test_csv_preserves_quoted_multiline_and_chinese_encoding():
    data = '姓名,说明\r\n学姐,"第一行\n第二行"\r\n兰琛,中文资料\r\n'.encode("gb18030")
    validate_csv(data)
    parsed = parse_csv(data)
    assert "学姐" in parsed.plain_text and "第一行\n第二行" in parsed.plain_text
    assert parsed.sections[0].anchor == {"row_start": 1, "row_end": 3}
    assert "line_start" not in parsed.sections[0].anchor
    assert_source_offsets(parsed)


def test_csv_rejects_unterminated_quoted_record():
    with pytest.raises(ParseError, match="CSV"):
        validate_csv(b'column\n"unfinished')


def test_html_extracts_headings_tables_pre_and_excludes_active_hidden_content():
    data = b'''<!DOCTYPE html><html><head><title>Private title</title><script src="https://external.invalid/app.js"></script></head><body>
    <h1>Title</h1><p>Visible <b>evidence</b>.</p><script>alert("do-not-read")</script>
    <style>.hidden{display:none}</style><div hidden>secret hidden</div><p aria-hidden="true">aria secret</p>
    <div style="display:none">css secret</div><iframe src="file:///etc/passwd">frame secret</iframe>
    <h2>Table</h2><table><tr><th>Item</th><th>Count</th></tr><tr><td>Slots</td><td>28</td></tr></table>
    <pre>line one\nline two</pre><img src="https://external.invalid/pixel"/></body></html>'''
    parsed = parse_html(data)
    assert [section.heading for section in parsed.sections] == ["Title", "Table"]
    assert parsed.sections[1].parent_index == 0
    assert "Slots | 28" in parsed.plain_text
    assert "line one\nline two" in parsed.plain_text
    for excluded in ("Private title", "do-not-read", "secret", "external.invalid", "display:none"):
        assert excluded not in parsed.plain_text
    assert_source_offsets(parsed)


def test_html_fragment_and_htm_alias_preserve_text():
    parsed = parse_html("<h2>中文资料</h2><div>真实正文</div>".encode(), "htm")
    assert parsed.metadata["format"] == "htm"
    assert "真实正文" in parsed.plain_text


def test_json_array_retains_typed_values_and_record_ranges():
    data = json.dumps([{"标题": "中文", "数量": 28, "启用": True}, {"nested": {"value": None}}], ensure_ascii=False).encode()
    validate_json(data)
    parsed = parse_json(data)
    assert parsed.sections[0].anchor == {"record_start": 1, "record_end": 2}
    assert '"数量": 28' in parsed.plain_text and '"启用": true' in parsed.plain_text
    assert '"value": null' in parsed.plain_text
    assert_source_offsets(parsed)


@pytest.mark.parametrize("raw", [b"{broken", b'{"value": NaN}', b'{"value": Infinity}'])
def test_json_rejects_invalid_structure_and_nonfinite_values(raw):
    with pytest.raises(ParseError, match="JSON"):
        validate_json(raw)


def test_jsonl_uses_actual_physical_lines_even_with_blank_records():
    raw = b'{"title":"first"}\n\n{"title":"second"}\n'
    validate_jsonl(raw)
    parsed = parse_jsonl(raw)
    assert parsed.sections[0].anchor == {"line_start": 1, "line_end": 3, "record_start": 1, "record_end": 2}
    assert "first" in parsed.plain_text and "second" in parsed.plain_text
    assert_source_offsets(parsed)


def test_jsonl_reports_the_bad_source_line():
    with pytest.raises(ParseError, match="第 3 行"):
        parse_jsonl(b'{"ok":true}\n\nbad-json')


def test_epub_uses_spine_order_and_local_chapter_paths():
    data = epub_bytes()
    validate_epub(data)
    parsed = parse_epub(data)
    assert [section.heading for section in parsed.sections] == ["第二章", "第一章"]
    assert parsed.sections[0].anchor == {"chapter_index": 1, "chapter_path": "OPS/chapter-two.xhtml"}
    assert "不应被提取" not in parsed.plain_text
    assert "tracker.png" not in parsed.plain_text
    assert_source_offsets(parsed)


@pytest.mark.parametrize("arguments", [{"encrypted": True}, {"href": "../../outside.xhtml"}, {"href": "https://external.invalid/body.xhtml"}])
def test_epub_rejects_encryption_and_external_or_parent_paths(arguments):
    with pytest.raises(ParseError):
        validate_epub(epub_bytes(**arguments))


@pytest.mark.parametrize("validator,fixture", [(validate_pptx, presentation_bytes), (validate_xlsx, workbook_bytes), (validate_epub, epub_bytes)])
def test_archives_are_bounded_before_real_extraction(validator, fixture):
    with pytest.raises(ParseError, match="限制"):
        validator(fixture(), max_uncompressed_bytes=10)


@pytest.mark.parametrize("member", ["../escape.xml", "C:/outside.xml", "ppt/vbaProject.bin", "ppt/unsafe.xml"])
def test_ooxml_rejects_unsafe_members_and_xml_entities(member):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
        archive.writestr(member, '<!DOCTYPE root [<!ENTITY x "bomb">]><root>&x;</root>')
    with pytest.raises(ParseError):
        validate_pptx(stream.getvalue())


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-32"])
def test_xml_declarations_cannot_bypass_archive_screening_with_encoding(encoding):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
        archive.writestr("ppt/unsafe.xml", '<!DOCTYPE root [<!ENTITY x "bomb">]><root>&x;</root>'.encode(encoding))
    with pytest.raises(ParseError, match="XML"):
        validate_pptx(stream.getvalue())


def test_document_builder_enforces_output_budget(monkeypatch):
    from app.ingestion.formats import common

    monkeypatch.setattr(common, "MAX_TEXT_CHARS", 10)
    builder = DocumentBuilder("test")
    with pytest.raises(ParseError, match="解析限制"):
        builder.add("too much", "long text beyond budget", {})
