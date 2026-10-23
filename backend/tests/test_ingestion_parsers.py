import io
import zipfile

import pymupdf
import pytest
from docx import Document

from app.ingestion.chunking import chunk_sections, count_tokens, split_text
from app.ingestion.parsers import ParseError, decode_text, parse_document, validate_content


def test_markdown_hierarchy_and_exact_line_anchors(tmp_path):
    text = "前言\n# 机器学习\n正文第一行\n## RAG\n检索增强生成。\n"
    path = tmp_path / "论文.md"
    path.write_text(text, encoding="utf-8")
    parsed = parse_document(path)
    assert [s.heading for s in parsed.sections] == ["正文", "机器学习", "RAG"]
    assert parsed.sections[2].parent_index == 1
    assert parsed.sections[2].anchor == {"line_start": 4, "line_end": 5}
    for section in parsed.sections:
        assert text[section.source_start:section.source_end] == section.text


def test_txt_chinese_encoding_and_binary_rejection(tmp_path):
    path = tmp_path / "中文.txt"
    path.write_bytes("这是中文知识库\n检索测试".encode("gb18030"))
    parsed = parse_document(path)
    assert "中文知识库" in parsed.plain_text
    assert parsed.sections[0].anchor["line_start"] == 1
    assert decode_text("中文".encode("utf-16")) == "中文"
    with pytest.raises(ParseError, match="二进制"):
        validate_content(b"binary\x00payload", "txt")


def test_pdf_pages_toc_and_scan_error(tmp_path):
    path = tmp_path / "text.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Real evidence page one")
    document.new_page().insert_text((72, 72), "Second page evidence")
    document.set_toc([[1, "Methods", 1]])
    document.save(path)
    document.close()
    parsed = parse_document(path)
    assert len(parsed.sections) == 2
    assert parsed.sections[1].anchor == {"page_start": 2, "page_end": 2}
    assert parsed.sections[0].heading == "Methods"
    empty_path = tmp_path / "scan.pdf"
    empty = pymupdf.open()
    empty.new_page()
    empty.save(empty_path)
    empty.close()
    with pytest.raises(ParseError, match="不支持 OCR"):
        parse_document(empty_path)


def test_docx_preserves_table_order_and_has_no_fake_pages(tmp_path):
    path = tmp_path / "research.docx"
    document = Document()
    document.add_heading("研究", 1)
    document.add_paragraph("前面的证据")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表格数据"
    table.cell(0, 1).text = "结果"
    document.add_paragraph("后面的证据")
    document.save(path)
    parsed = parse_document(path)
    assert parsed.plain_text.index("前面的证据") < parsed.plain_text.index("表格数据") < parsed.plain_text.index("后面的证据")
    assert all("page_start" not in section.anchor for section in parsed.sections)
    assert parsed.sections[-1].anchor["paragraph_end"] == 4


def test_file_structure_validation():
    for data, ext in [(b"hello", "pdf"), (b"PKbroken", "docx"), (b"text", "exe"), (b"", "txt")]:
        with pytest.raises(ParseError):
            validate_content(data, ext)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("word/vbaProject.bin", "macro")
    with pytest.raises(ParseError, match="宏"):
        validate_content(output.getvalue(), "docx")


def test_docx_merged_cell_is_once_per_row_but_equal_independent_values_remain(tmp_path):
    path = tmp_path / "merged-form.docx"
    document = Document()
    document.add_paragraph("资格申请表")
    table = document.add_table(rows=2, cols=7)
    merged = table.cell(0, 0).merge(table.cell(0, 6))
    merged.text = "信息与智能科学学院"
    table.cell(1, 0).text = "相同内容"
    table.cell(1, 1).text = "相同内容"
    document.save(path)
    parsed = parse_document(path)
    assert parsed.plain_text.count("信息与智能科学学院") == 1
    assert parsed.plain_text.count("相同内容") == 2
    assert "相同内容 | 相同内容" in parsed.plain_text
    assert parsed.sections[0].anchor == {"paragraph_start": 1, "paragraph_end": 3}


def test_chunking_is_deterministic_bounded_and_covers_source(tmp_path):
    path = tmp_path / "long.md"
    text = "# 长文\n" + "这是需要保留原始定位的研究证据，检索时必须能找到。\n" * 150
    path.write_text(text, encoding="utf-8")
    parsed = parse_document(path)
    chunks = chunk_sections(parsed.sections, size=100, overlap=15)
    assert chunks == chunk_sections(parsed.sections, size=100, overlap=15)
    assert len(chunks) > 10
    assert all(chunk.token_count <= 100 for chunk in chunks)
    for chunk in chunks:
        assert text[chunk.source_start:chunk.source_end] == chunk.content
        assert chunk.anchor["line_start"] >= 1
    assert chunks[1].source_start < chunks[0].source_end
    assert chunks[-1].source_end == len(text)
    with pytest.raises(ValueError):
        split_text(text, 10, 10)
    assert count_tokens("知识库") > 0
