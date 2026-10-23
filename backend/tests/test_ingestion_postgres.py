"""Opt-in real PostgreSQL/pgvector checks. Never use the application database.

PG_INGESTION_TEST_URL must end in /zhicluster_ingestion_test.
The database must exist; Alembic initializes it. Tests retain inspectable evidence.
"""

import hashlib
import json
import os
import subprocess
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from uuid import uuid4

import httpx
import pymupdf
import pytest
from docx import Document as WordDocument
from sqlalchemy import event, func, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import sessionmaker

from app.ai.gateway import ModelGateway
from app.core.config import Settings
from app.core.errors import AppError
from app.db.models import (
    ChunkEmbedding,
    Document,
    DocumentChunk,
    DocumentSection,
    EmbeddingProfile,
    IngestionJob,
    KnowledgeBase,
    KnowledgeBaseMember,
    User,
    utcnow,
)
from app.db.session import make_engine
from app.retrieval.search import search_knowledge
from app.worker.pipeline import process_ingestion_job
from app.worker.queue import claim_job, heartbeat

PG_URL = os.getenv("PG_INGESTION_TEST_URL", "")
pytestmark = pytest.mark.skipif(not PG_URL, reason="Set PG_INGESTION_TEST_URL to isolated PostgreSQL ingestion test database")


@pytest.fixture(scope="module")
def pg_factory():
    assert make_url(PG_URL).database == "zhicluster_ingestion_test", "Refusing to migrate a non-test database"
    backend = Path(__file__).resolve().parents[1]
    environment = {key: os.environ[key] for key in ("PATH", "SYSTEMROOT", "USERPROFILE", "TEMP", "TMP") if key in os.environ}
    environment.update({"DATABASE_URL": PG_URL, "APP_ENV": "test", "PYTHONIOENCODING": "utf-8"})
    migration = subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=backend, env=environment, check=False, capture_output=True, text=True, encoding="utf-8")
    assert migration.returncode == 0, migration.stdout + migration.stderr
    engine = make_engine(PG_URL)
    factory = sessionmaker(engine, expire_on_commit=False)
    yield factory
    engine.dispose()


def create_input(factory, root, ext="md", content=None):
    suffix = uuid4().hex[:12]
    if content is None:
        content = "# 机器学习研究\n这是中文真实证据。检索增强生成与知识库范围。\n".encode()
    with factory() as db:
        actor = User(email=f"ingestion-{suffix}@example.test", display_name="入库集成测试")
        db.add(actor)
        db.flush()
        kb = KnowledgeBase(owner_id=actor.id, name=f"入库验收-{suffix}")
        db.add(kb)
        db.flush()
        db.add(KnowledgeBaseMember(kb_id=kb.id, user_id=actor.id, role="owner"))
        document = Document(kb_id=kb.id, owner_id=actor.id, display_name=f"中文证据.{ext}", original_filename=f"中文证据.{ext}", mime_type="application/octet-stream", ext=ext, sha256=hashlib.sha256(content).hexdigest(), size_bytes=len(content), storage_key=f"{kb.id}/{uuid4()}/original")
        db.add(document)
        db.flush()
        path = root / document.storage_key
        path.parent.mkdir(parents=True)
        path.write_bytes(content)
        job = IngestionJob(document_id=document.id, kb_id=kb.id)
        db.add(job)
        db.commit()
        return actor.id, kb.id, document.id, job.id


def config(root, **overrides):
    return Settings(_env_file=None, upload_root=root, database_url=PG_URL, embedding_dimension=3, chunk_size_tokens=100, chunk_overlap_tokens=10, **overrides)


def test_pg_skip_locked_exclusive_claim_and_lease_recovery(pg_factory, tmp_path):
    _, _, _, first_id = create_input(pg_factory, tmp_path)
    _, _, _, second_id = create_input(pg_factory, tmp_path)
    # Ignore prior completed evidence and isolate the queue candidates created here.
    with pg_factory() as locked_db, pg_factory() as claimant_db:
        locked_db.execute(select(IngestionJob).where(IngestionJob.id == first_id).with_for_update()).scalar_one()
        claimed = claim_job(claimant_db, "pg-worker-b", 60)
        assert claimed.id == second_id
        assert heartbeat(claimant_db, second_id, "wrong-worker", 60) is False
        assert heartbeat(claimant_db, second_id, "pg-worker-b", 60) is True
        assert claim_job(claimant_db, "pg-worker-c", 60) is None
        locked_db.rollback()
        claimed_first = claim_job(claimant_db, "pg-worker-a", 60)
        assert claimed_first.id == first_id
        assert claimed_first.id != claimed.id
        claimed_first.lease_until = utcnow() - timedelta(seconds=5)
        claimant_db.commit()
    with pg_factory() as db:
        recovered = claim_job(db, "pg-worker-recovered", 60)
        assert recovered.id == first_id
        assert recovered.attempts == 2
        assert recovered.lease_owner == "pg-worker-recovered"
        for job_id in [first_id, second_id]:
            db.get(IngestionJob, job_id).status = "cancelled"
        db.commit()


@pytest.mark.parametrize("ext", ["pdf", "docx", "md", "txt"])
def test_pg_real_four_format_pipeline_and_keyword_search(pg_factory, tmp_path, ext):
    if ext == "pdf":
        document = pymupdf.open()
        document.new_page().insert_text((72, 72), "机器学习研究与中文证据", fontname="china-s")
        document.new_page().insert_text((72, 72), "第二页检索增强生成", fontname="china-s")
        content = document.tobytes()
        document.close()
    elif ext == "docx":
        path = tmp_path / "fixture.docx"
        document = WordDocument()
        document.add_heading("机器学习", 1)
        document.add_paragraph("这是中文证据，DOCX不能虚构页码。")
        document.save(path)
        content = path.read_bytes()
    else:
        content = "# 机器学习\n这是中文证据，真实文档内容必须可检索。\n".encode("utf-8")
    actor_id, kb_id, document_id, job_id = create_input(pg_factory, tmp_path / "uploads", ext, content)
    settings = config(tmp_path / "uploads")
    process_ingestion_job(job_id, pg_factory, settings)
    with pg_factory() as db:
        stored = db.get(Document, document_id)
        assert stored.parse_status == "ready"
        assert db.get(IngestionJob, job_id).status == "partial_ready"
        sections = db.scalars(select(DocumentSection).where(DocumentSection.document_id == document_id)).all()
        assert sections
        if ext == "pdf":
            assert {section.anchors["page_start"] for section in sections} == {1, 2}
        elif ext == "docx":
            assert all("paragraph_start" in section.anchors and "page_start" not in section.anchors for section in sections)
        else:
            assert sections[0].anchors["line_start"] == 1
        results = search_knowledge(db, db.get(User, actor_id), {"query": "中文证据", "kb_ids": [kb_id]}, settings)
        assert results["mode"] == "keyword"
        assert results["items"]
        assert {item["document_id"] for item in results["items"]} == {document_id}


def test_pg_empty_pdf_explicitly_fails(pg_factory, tmp_path):
    pdf = pymupdf.open()
    pdf.new_page()
    content = pdf.tobytes()
    pdf.close()
    _, _, document_id, job_id = create_input(pg_factory, tmp_path, "pdf", content)
    process_ingestion_job(job_id, pg_factory, config(tmp_path))
    with pg_factory() as db:
        assert db.get(Document, document_id).parse_status == "failed"
        assert "不支持 OCR" in db.get(IngestionJob, job_id).error


def test_pg_profile_switch_exact_cosine_and_scope_filter(pg_factory, tmp_path, monkeypatch):
    actor_id, kb_id, document_id, job_id = create_input(pg_factory, tmp_path)
    settings = config(tmp_path, embedding_base_url="https://test-only.invalid/v1", embedding_api_key="test-only", embedding_model=f"deterministic-test-{uuid4().hex[:8]}")

    class FakeGateway(ModelGateway):
        def embed(self, texts):
            return [[1.0] + [0.1] * (self.settings.embedding_dimension - 1) for _ in texts]

    process_ingestion_job(job_id, pg_factory, settings, FakeGateway(settings))
    old_fingerprint = FakeGateway(settings).profile_fingerprint
    settings.embedding_dimension = 4
    with pg_factory() as db:
        reindex = IngestionJob(document_id=document_id, kb_id=kb_id, kind="reindex")
        db.add(reindex)
        db.commit()
        reindex_id = reindex.id
    process_ingestion_job(reindex_id, pg_factory, settings, FakeGateway(settings))
    new_fingerprint = FakeGateway(settings).profile_fingerprint
    assert old_fingerprint != new_fingerprint
    captured = []

    def record_sql(connection, cursor, statement, parameters, context, executemany):
        captured.append(statement)

    event.listen(pg_factory.kw["bind"], "before_cursor_execute", record_sql)
    monkeypatch.setattr("app.retrieval.search.ModelGateway", FakeGateway)
    try:
        with pg_factory() as db:
            profiles = db.scalars(select(EmbeddingProfile).where(EmbeddingProfile.fingerprint.in_([old_fingerprint, new_fingerprint]))).all()
            assert {profile.dimension for profile in profiles} == {3, 4}
            chunks = db.scalar(select(func.count(DocumentChunk.id)).where(DocumentChunk.document_id == document_id))
            counts = [db.scalar(select(func.count(ChunkEmbedding.id)).where(ChunkEmbedding.profile_id == profile.id)) for profile in profiles]
            assert counts == [chunks, chunks]
            actor = db.get(User, actor_id)
            results = search_knowledge(db, actor, {"query": "语义测试无关键词匹配", "kb_ids": [kb_id], "mode": "hybrid"}, settings)
            assert results["mode"] == "hybrid"
            assert results["items"]
            assert {item["document_id"] for item in results["items"]} == {document_id}
            assert any("<=>" in sql and "document_chunks.kb_id IN" in sql and "chunk_embeddings.profile_id =" in sql for sql in captured)
            foreign_kb = db.scalar(select(KnowledgeBase).where(KnowledgeBase.id != kb_id))
            assert foreign_kb is not None
            with pytest.raises(AppError):
                search_knowledge(db, actor, {"query": "中文", "kb_ids": [foreign_kb.id]}, settings)
    finally:
        event.remove(pg_factory.kw["bind"], "before_cursor_execute", record_sql)


def test_pg_old_worker_cannot_save_vectors_after_lease_transfer(pg_factory, tmp_path):
    _, _, document_id, job_id = create_input(pg_factory, tmp_path)
    settings = config(tmp_path, embedding_base_url="https://test-only.invalid/v1", embedding_api_key="test-only", embedding_model=f"lease-test-{uuid4().hex[:8]}")
    with pg_factory() as db:
        job = db.get(IngestionJob, job_id)
        job.status = "running"
        job.lease_owner = "old-worker"
        job.lease_until = utcnow() + timedelta(seconds=60)
        db.commit()

    class TransferLeaseGateway(ModelGateway):
        def embed(self, texts):
            with pg_factory() as db:
                db.get(IngestionJob, job_id).lease_owner = "new-worker"
                db.commit()
            return [[1.0, 0.1, 0.2] for _ in texts]

    process_ingestion_job(job_id, pg_factory, settings, TransferLeaseGateway(settings), "old-worker")
    with pg_factory() as db:
        assert db.get(IngestionJob, job_id).lease_owner == "new-worker"
        chunk_ids = select(DocumentChunk.id).where(DocumentChunk.document_id == document_id)
        assert db.scalar(select(func.count(ChunkEmbedding.id)).where(ChunkEmbedding.chunk_id.in_(chunk_ids))) == 0
        assert db.get(Document, document_id).index_status != "ready"
        db.get(IngestionJob, job_id).status = "cancelled"
        db.commit()


def test_pg_parallel_documents_share_profile_and_overlap_provider_requests(pg_factory, tmp_path):
    inputs = [create_input(pg_factory, tmp_path) for _ in range(2)]
    settings = config(tmp_path, embedding_base_url="https://parallel.invalid/v1", embedding_api_key="test-only", embedding_model=f"parallel-{uuid4().hex[:8]}")
    barrier = threading.Barrier(2, timeout=8)
    lock = threading.Lock()
    overlapping = []
    active = 0

    def respond(request):
        nonlocal active
        payload = json.loads(request.content)
        with lock:
            active += 1
            overlapping.append(active)
        try:
            barrier.wait()
            return httpx.Response(200, json={"data": [{"index": index, "embedding": [1.0, 0.1, 0.2]} for index, _ in enumerate(payload["input"])]})
        finally:
            with lock:
                active -= 1

    with pg_factory() as db:
        for _, _, _, job_id in inputs:
            job = db.get(IngestionJob, job_id)
            job.status, job.lease_owner = "running", f"parallel-{job_id}"
            job.lease_until = utcnow() + timedelta(seconds=60)
        db.commit()
    transport = httpx.MockTransport(respond)
    with ThreadPoolExecutor(max_workers=2) as pool:
        tasks = [pool.submit(process_ingestion_job, job_id, pg_factory, settings, ModelGateway(settings, transport=transport, workload="background"), f"parallel-{job_id}") for _, _, _, job_id in inputs]
        for task in tasks:
            task.result(timeout=15)
    assert max(overlapping) == 2
    with pg_factory() as db:
        profiles = list(db.scalars(select(EmbeddingProfile).where(EmbeddingProfile.fingerprint == ModelGateway(settings).profile_fingerprint)))
        assert len(profiles) == 1
        for _, _, document_id, job_id in inputs:
            assert db.get(Document, document_id).index_status == "ready"
            assert db.get(IngestionJob, job_id).status == "partial_ready"


def test_pg_metadata_edit_during_parse_uses_same_namespace_order(pg_factory, tmp_path, monkeypatch):
    from app.worker import pipeline

    _, kb_id, document_id, job_id = create_input(pg_factory, tmp_path)
    settings = config(tmp_path)
    namespace_held = threading.Event()
    worker_waiting_namespace = threading.Event()
    parsed_ready = threading.Event()
    main_thread = threading.get_ident()
    original_parser = pipeline.parse_document
    metadata_tasks = []

    def edit_metadata():
        with pg_factory() as db:
            db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id).with_for_update())
            namespace_held.set()
            assert worker_waiting_namespace.wait(4), "pipeline took a document write lock before the namespace"
            document = db.scalar(select(Document).where(Document.id == document_id).with_for_update())
            document.display_name = "concurrent-metadata.txt"
            document.version += 1
            db.commit()

    def observe(connection, cursor, statement, parameters, context, executemany):
        if threading.get_ident() == main_thread and parsed_ready.is_set() and "FROM knowledge_bases" in statement and "FOR UPDATE" in statement:
            worker_waiting_namespace.set()

    with ThreadPoolExecutor(max_workers=1) as editor:
        def parse_then_edit(*args):
            parsed = original_parser(*args)
            metadata_tasks.append(editor.submit(edit_metadata))
            assert namespace_held.wait(3)
            parsed_ready.set()
            return parsed

        monkeypatch.setattr(pipeline, "parse_document", parse_then_edit)
        event.listen(pg_factory.kw["bind"], "before_cursor_execute", observe)
        try:
            process_ingestion_job(job_id, pg_factory, settings)
            metadata_tasks[0].result(timeout=6)
        finally:
            event.remove(pg_factory.kw["bind"], "before_cursor_execute", observe)
    with pg_factory() as db:
        document = db.get(Document, document_id)
        assert document.display_name == "concurrent-metadata.txt"
        assert document.parse_status == "ready"
        assert db.get(IngestionJob, job_id).status == "partial_ready"


def test_pg_claim_rechecks_sibling_after_candidate_snapshot_before_document_lock(pg_factory, tmp_path):
    _, kb_id, document_id, ingest_id = create_input(pg_factory, tmp_path)
    with pg_factory() as db:
        reindex = IngestionJob(kb_id=kb_id, document_id=document_id, kind="reindex")
        db.add(reindex)
        db.commit()
        reindex_id = reindex.id
    candidate_read = threading.Event()
    resume_first = threading.Event()
    first_thread = []

    def pause_after_candidates(connection, cursor, statement, parameters, context, executemany):
        if first_thread and threading.get_ident() == first_thread[0] and not candidate_read.is_set() and "FROM ingestion_jobs" in statement and "FOR UPDATE SKIP LOCKED" in statement:
            candidate_read.set()
            assert resume_first.wait(6), "test did not resume the first candidate snapshot"

    def first_claim():
        first_thread.append(threading.get_ident())
        with pg_factory() as db:
            claimed = claim_job(db, "candidate-before-sibling", 60, kinds=("ingest",))
            return claimed.id if claimed else None

    event.listen(pg_factory.kw["bind"], "after_cursor_execute", pause_after_candidates)
    try:
        with ThreadPoolExecutor(max_workers=1) as pool:
            first = pool.submit(first_claim)
            try:
                assert candidate_read.wait(3)
                with pg_factory() as db:
                    # This worker did not see the first job running yet, and
                    # wins the document lock before the first worker reaches it.
                    winner = claim_job(db, "sibling-committed-first", 60, kinds=("reindex",))
                    assert winner.id == reindex_id
                resume_first.set()
                assert first.result(timeout=5) is None
            finally:
                resume_first.set()
        with pg_factory() as db:
            assert db.get(IngestionJob, ingest_id).status == "queued"
            assert db.get(IngestionJob, ingest_id).attempts == 0
            assert db.get(IngestionJob, reindex_id).status == "running"
    finally:
        event.remove(pg_factory.kw["bind"], "after_cursor_execute", pause_after_candidates)
        with pg_factory() as db:
            for job_id in (ingest_id, reindex_id):
                job = db.get(IngestionJob, job_id)
                job.status = "cancelled"
                job.lease_owner = job.lease_until = None
            db.commit()
