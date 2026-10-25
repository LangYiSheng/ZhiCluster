"""PDF and DOCX adapters retain the original extraction and anchor behavior."""

import io
import re
import statistics

import pymupdf
from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.parser_types import ParsedDocument, ParsedSection, ParseError


def validate_pdf(data, max_uncompressed_bytes=200 * 1024 * 1024):
    if not data.startswith(b"%PDF-"):
        raise ParseError("文件内容与 PDF 扩展名不符")


def validate_docx(data, max_uncompressed_bytes=200 * 1024 * 1024):
    from app.ingestion.formats.common import checked_archive

    with checked_archive(data, "DOCX", {"word/document.xml", "[Content_Types].xml"}, max_uncompressed_bytes):
        pass


def parse_pdf(data: bytes) -> ParsedDocument:
    sections = []
    texts = []
    offset = 0
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ParseError("PDF 已加密，请先移除密码")
            toc = document.get_toc()
            for page_index, page in enumerate(document):
                text = page.get_text("text", sort=True)
                texts.append(text)
                if text.strip():
                    candidates = [entry for entry in toc if entry[2] <= page_index + 1]
                    heading = candidates[-1][1] if candidates else _pdf_heading(page) or f"第 {page_index + 1} 页"
                    level = int(candidates[-1][0]) if candidates else 1
                    sections.append(ParsedSection(heading, level, text, {"page_start": page_index + 1, "page_end": page_index + 1}, offset, offset + len(text)))
                offset += len(text) + 1
            if not sections:
                raise ParseError("缺少文本层，首版不支持 OCR")
            return ParsedDocument(sections, "\n".join(texts), {"pages": len(document), "format": "pdf"})
    except (pymupdf.FileDataError, RuntimeError) as error:
        raise ParseError("PDF 无法读取或文件已损坏") from error


def _pdf_heading(page) -> str | None:
    spans = [span for block in page.get_text("dict").get("blocks", []) if block.get("type") == 0 for line in block.get("lines", []) for span in line.get("spans", []) if span.get("text", "").strip()]
    if not spans:
        return None
    typical_size = statistics.median(span.get("size", 0) for span in spans)
    titles = [span for span in spans if span.get("size", 0) >= typical_size * 1.2 and len(span.get("text", "")) <= 120]
    return titles[0]["text"].strip() if titles else None


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = WordDocument(io.BytesIO(data))
    except Exception as error:
        raise ParseError("DOCX 无法读取或文件已损坏") from error
    sections: list[ParsedSection] = []
    pieces = []
    blocks: list[tuple[str, int]] = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            style = paragraph.style.name if paragraph.style else ""
            match = re.search(r"(?:Heading|标题)\s*(\d+)", style, re.IGNORECASE)
            blocks.append((paragraph.text, int(match.group(1)) if match else 0))
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            for row in table.rows:
                # python-docx repeats the same XML cell for each occupied grid
                # column of a horizontal merge. Keep one copy of that cell,
                # while preserving equal text in distinct, meaningful cells.
                seen_cells = set()
                texts = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        texts.append(cell.text)
                blocks.append((" | ".join(texts), 0))
    starts = []
    offset = 0
    for index, (text, level) in enumerate(blocks):
        pieces.append(text)
        if level:
            starts.append((index, offset, text, level))
        offset += len(text) + 1
    full_text = "\n".join(pieces)
    if not starts or starts[0][0] > 0:
        starts.insert(0, (0, 0, "正文", 0))
    stack = []
    for position, (paragraph_index, start, heading, level) in enumerate(starts):
        paragraph_end = starts[position + 1][0] if position + 1 < len(starts) else len(blocks)
        end = starts[position + 1][1] if position + 1 < len(starts) else len(full_text)
        text = full_text[start:end]
        if not text.strip():
            continue
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else None
        sections.append(ParsedSection(heading, level, text, {"paragraph_start": paragraph_index + 1, "paragraph_end": paragraph_end}, start, end, parent))
        stack.append((level, len(sections) - 1))
    return ParsedDocument(sections, full_text, {"paragraphs": len(blocks), "format": "docx"})
