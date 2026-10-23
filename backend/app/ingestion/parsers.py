"""Extract real text and honest source anchors; never invent document page numbers."""

from __future__ import annotations

import io
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from charset_normalizer import from_bytes
from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.structure import FENCE


class ParseError(ValueError):
    """A user-correctable file format or document content error."""


@dataclass(frozen=True)
class ParsedSection:
    heading: str
    level: int
    text: str
    anchor: dict[str, int]
    source_start: int
    source_end: int
    parent_index: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    sections: list[ParsedSection]
    plain_text: str
    metadata: dict = field(default_factory=dict)


SUPPORTED_EXTENSIONS = {"pdf", "docx", "md", "txt"}


def validate_content(data: bytes, ext: str, max_uncompressed_bytes: int = 200 * 1024 * 1024) -> None:
    ext = ext.lower().lstrip(".")
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError("仅支持 PDF、DOCX、Markdown 和 TXT 文件")
    if not data:
        raise ParseError("文件为空")
    if ext == "pdf" and not data.startswith(b"%PDF-"):
        raise ParseError("文件内容与 PDF 扩展名不符")
    if ext == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if sum(m.file_size for m in members) > max_uncompressed_bytes:
                    raise ParseError("DOCX 解压后的内容超过安全限制")
                names = {m.filename for m in members}
                if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                    raise ParseError("文件不是有效的 DOCX 文档")
                if any("vbaproject" in name.lower() for name in names):
                    raise ParseError("首版不接收含宏的文档")
        except zipfile.BadZipFile as error:
            raise ParseError("文件内容与 DOCX 扩展名不符") from error
    if ext in {"md", "txt"}:
        decode_text(data)


def decode_text(data: bytes) -> str:
    if b"\x00" in data and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ParseError("文本文件包含二进制内容")
    for encoding in ("utf-8-sig", "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "gb18030"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text.replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeError:
            pass
    guess = from_bytes(data).best()
    if guess is None or guess.chaos > 0.3:
        raise ParseError("无法识别文本编码，请转换为 UTF-8 后重试")
    return str(guess).replace("\r\n", "\n").replace("\r", "\n")


def _text_sections(text: str, markdown: bool) -> list[ParsedSection]:
    lines = text.splitlines(keepends=True)
    starts = []
    offset = 0
    fence = None
    for index, line in enumerate(lines):
        marker = FENCE.match(line.rstrip("\n")) if markdown else None
        inside_fence = bool(fence)
        if marker:
            if fence and marker[1][0] == fence[0] and len(marker[1]) >= len(fence) and not marker[2].strip():
                fence = None
            elif not fence:
                fence = marker[1]
        heading = re.match(r"^ {0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line.rstrip("\n")) if markdown and not inside_fence and not marker else None
        if heading:
            starts.append((index, offset, heading.group(2), len(heading.group(1))))
        offset += len(line)
    if not starts or starts[0][0] > 0:
        starts.insert(0, (0, 0, "正文", 0))
    sections: list[ParsedSection] = []
    stack: list[tuple[int, int]] = []
    for position, (line_index, start, heading, level) in enumerate(starts):
        next_line = starts[position + 1][0] if position + 1 < len(starts) else len(lines)
        end = starts[position + 1][1] if position + 1 < len(starts) else len(text)
        body = text[start:end]
        if not body.strip():
            continue
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else None
        section = ParsedSection(heading, level, body, {"line_start": line_index + 1, "line_end": next_line}, start, end, parent)
        stack.append((level, len(sections)))
        sections.append(section)
    return sections


def _parse_pdf(data: bytes) -> ParsedDocument:
    sections = []
    texts = []
    offset = 0
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ParseError("PDF 已加密，请先移除密码")
            toc = document.get_toc()
            for page_index, page in enumerate(document):
                text = page.get_text("text", sort=True)
                texts.append(text)
                if text.strip():
                    candidates = [entry for entry in toc if entry[2] <= page_index + 1]
                    heading = candidates[-1][1] if candidates else _pdf_heading(page) or f"第 {page_index + 1} 页"
                    level = int(candidates[-1][0]) if candidates else 1
                    sections.append(ParsedSection(heading, level, text, {"page_start": page_index + 1, "page_end": page_index + 1}, offset, offset + len(text)))
                offset += len(text) + 1
            if not sections:
                raise ParseError("缺少文本层，首版不支持 OCR")
            return ParsedDocument(sections, "\n".join(texts), {"pages": len(document), "format": "pdf"})
    except (pymupdf.FileDataError, RuntimeError) as error:
        raise ParseError("PDF 无法读取或文件已损坏") from error


def _pdf_heading(page) -> str | None:
    spans = [span for block in page.get_text("dict").get("blocks", []) if block.get("type") == 0 for line in block.get("lines", []) for span in line.get("spans", []) if span.get("text", "").strip()]
    if not spans:
        return None
    typical_size = statistics.median(span.get("size", 0) for span in spans)
    titles = [span for span in spans if span.get("size", 0) >= typical_size * 1.2 and len(span.get("text", "")) <= 120]
    return titles[0]["text"].strip() if titles else None


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = WordDocument(io.BytesIO(data))
    except Exception as error:
        raise ParseError("DOCX 无法读取或文件已损坏") from error
    sections: list[ParsedSection] = []
    pieces = []
    blocks: list[tuple[str, int]] = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            style = paragraph.style.name if paragraph.style else ""
            match = re.search(r"(?:Heading|标题)\s*(\d+)", style, re.IGNORECASE)
            blocks.append((paragraph.text, int(match.group(1)) if match else 0))
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            for row in table.rows:
                # python-docx repeats the same XML cell for each occupied grid
                # column of a horizontal merge. Keep one copy of that cell,
                # while preserving equal text in distinct, meaningful cells.
                seen_cells = set()
                texts = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        texts.append(cell.text)
                blocks.append((" | ".join(texts), 0))
    starts = []
    offset = 0
    for index, (text, level) in enumerate(blocks):
        pieces.append(text)
        if level:
            starts.append((index, offset, text, level))
        offset += len(text) + 1
    full_text = "\n".join(pieces)
    if not starts or starts[0][0] > 0:
        starts.insert(0, (0, 0, "正文", 0))
    stack = []
    for position, (paragraph_index, start, heading, level) in enumerate(starts):
        paragraph_end = starts[position + 1][0] if position + 1 < len(starts) else len(blocks)
        end = starts[position + 1][1] if position + 1 < len(starts) else len(full_text)
        text = full_text[start:end]
        if not text.strip():
            continue
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else None
        sections.append(ParsedSection(heading, level, text, {"paragraph_start": paragraph_index + 1, "paragraph_end": paragraph_end}, start, end, parent))
        stack.append((level, len(sections) - 1))
    return ParsedDocument(sections, full_text, {"paragraphs": len(blocks), "format": "docx"})


def parse_document(path: str | Path, ext: str | None = None) -> ParsedDocument:
    file_path = Path(path)
    extension = (ext or file_path.suffix).lower().lstrip(".")
    data = file_path.read_bytes()
    validate_content(data, extension)
    if extension == "pdf":
        parsed = _parse_pdf(data)
    elif extension == "docx":
        parsed = _parse_docx(data)
    else:
        text = decode_text(data)
        parsed = ParsedDocument(_text_sections(text, extension == "md"), text, {"format": extension, "lines": len(text.splitlines())})
    if not parsed.sections:
        raise ParseError("文档没有可提取的正文")
    return parsed
