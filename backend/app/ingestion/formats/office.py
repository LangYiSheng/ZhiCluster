"""Read OOXML slides and workbooks without running Office or formulas."""

from __future__ import annotations

import io
from datetime import date, datetime, time

from app.ingestion.formats.common import DEFAULT_ARCHIVE_LIMIT, DocumentBuilder, checked_archive
from app.ingestion.parser_types import ParseError

MAX_SHEETS = 256
MAX_ROWS = 100_000
MAX_COLUMNS = 1_024
MAX_CELLS = 1_000_000


def validate_pptx(data: bytes, max_uncompressed_bytes: int = DEFAULT_ARCHIVE_LIMIT):
    with checked_archive(data, "PPTX", {"[Content_Types].xml", "ppt/presentation.xml"}, max_uncompressed_bytes):
        pass


def validate_xlsx(data: bytes, max_uncompressed_bytes: int = DEFAULT_ARCHIVE_LIMIT):
    with checked_archive(data, "XLSX", {"[Content_Types].xml", "xl/workbook.xml"}, max_uncompressed_bytes):
        pass


def _shape_text(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from _shape_text(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                seen = set()
                cells = []
                for cell in row.cells:
                    if not cell.is_spanned and cell._tc not in seen:
                        seen.add(cell._tc)
                        cells.append(cell.text)
                yield " | ".join(cells)
        elif getattr(shape, "has_text_frame", False):
            yield shape.text_frame.text


def parse_pptx(data: bytes):
    try:
        from pptx import Presentation
    except ImportError as error:
        raise ParseError("PPTX 解析依赖未安装，请安装 python-pptx") from error
    validate_pptx(data)
    try:
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > 10_000:
            raise ParseError("幻灯片数量超过解析限制，请拆分文件")
        builder = DocumentBuilder("pptx")
        for index, slide in enumerate(presentation.slides, 1):
            pieces = list(_shape_text(slide.shapes))
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    pieces.append("备注：\n" + frame.text)
            title = slide.shapes.title
            heading = title.text.strip() if title is not None else ""
            builder.add(heading or f"幻灯片 {index}", "\n".join(pieces), {"slide_start": index, "slide_end": index})
        return builder.build(slides=len(presentation.slides), preview="extracted_text", image_text="not_ocr")
    except ParseError:
        raise
    except Exception as error:
        raise ParseError("PPTX 无法读取或文件已损坏") from error


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_xlsx(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ParseError("XLSX 解析依赖未安装，请安装 openpyxl") from error
    validate_xlsx(data)
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ParseError("工作表数量超过解析限制，请拆分文件")
        builder = DocumentBuilder("xlsx")
        scanned = 0
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            rows = []
            start_row = 1
            last_row = 0
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ParseError(f"工作表 {sheet.title} 超过 {MAX_COLUMNS} 列，请拆分文件")
            # Read-only worksheets stream their XML. Explicit limits also bound
            # incorrectly declared dimensions and sparsely formatted workbooks.
            sheet.reset_dimensions()
            for row_index, values in enumerate(sheet.iter_rows(values_only=True), 1):
                scanned += len(values)
                if row_index > MAX_ROWS or len(values) > MAX_COLUMNS or scanned > MAX_CELLS:
                    raise ParseError("工作簿行数或单元格数量超过解析限制，请拆分文件")
                if not any(value is not None for value in values):
                    continue
                rendered = [_cell_text(value) for value in values]
                while rendered and not rendered[-1]:
                    rendered.pop()
                if not rows:
                    start_row = row_index
                rows.append(f"行 {row_index}: " + " | ".join(rendered))
                last_row = row_index
                if len(rows) >= 100:
                    builder.add(f"{sheet.title} · 行 {start_row}–{last_row}", "\n".join(rows),
                        {"sheet_index": sheet_index, "sheet_name": sheet.title, "row_start": start_row, "row_end": last_row})
                    rows = []
            if rows:
                builder.add(f"{sheet.title} · 行 {start_row}–{last_row}", "\n".join(rows),
                    {"sheet_index": sheet_index, "sheet_name": sheet.title, "row_start": start_row, "row_end": last_row})
        return builder.build(sheets=len(workbook.worksheets), formulas="stored_text_not_evaluated", preview="extracted_text")
    except ParseError:
        raise
    except Exception as error:
        raise ParseError("XLSX 无法读取或文件已损坏") from error
    finally:
        if workbook is not None:
            workbook.close()
