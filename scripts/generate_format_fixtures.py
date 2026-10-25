"""Generate genuine document-format samples for isolated ingestion acceptance."""

from __future__ import annotations

import argparse
import csv
import io
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def generate(destination: Path, *, legacy: bool = False) -> list[dict]:
    destination.mkdir(parents=True, exist_ok=True)
    proof = "formatproof 课程申请名额为28人。"
    (destination / "sample.txt").write_text(proof + " TXT资料", encoding="utf-8")
    (destination / "sample.md").write_text("# Markdown资料\n\n" + proof, encoding="utf-8")
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Cover page")
        pdf.new_page().insert_text((72, 72), "formatproof PDF quota is 28 people.")
        pdf.save(destination / "sample.pdf")
    word = Document()
    word.add_heading("申请资料", 1)
    word.add_paragraph(proof + " DOCX资料")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "课程", "名额"
    table.cell(1, 0).text, table.cell(1, 1).text = "计算机", "28"
    word.save(destination / "sample.docx")
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = "申请资料封面"
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "第2张申请幻灯片"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2)).text_frame.text = proof + " PPTX资料"
    slides.save(destination / "sample.pptx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算表"
    sheet.append(["项目", "名额"])
    sheet.append([proof + " XLSX资料", 28])
    sheet.append(["仅保留公式", "=SUM(B2:B2)"])
    workbook.create_sheet("第二张表").append(["不同工作表的第1行", 9])
    workbook.save(destination / "sample.xlsx")
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["课程", "备注", "名额"])
    writer.writerow(["计算机", proof + ' CSV资料，含"引号"\n及换行', 28])
    (destination / "sample.csv").write_text(stream.getvalue(), encoding="utf-8-sig", newline="")
    for extension in ("html", "htm"):
        (destination / f"sample.{extension}").write_text(
            '<!doctype html><html><head><title>申请资料</title><style>.hidden{color:red}</style></head>'
            '<body><nav>导航文本</nav><h1>申请材料</h1><p>' + proof + f' {extension.upper()}资料</p>'
            '<script>window.__formatScriptExecuted=true;fetch("https://example.invalid/should-not-load")</script>'
            '<img src="https://example.invalid/should-not-load.png"></body></html>', encoding="utf-8")
    (destination / "sample.json").write_text(json.dumps({"申请记录": [{"正文": proof + " JSON资料", "名额": 28}]}, ensure_ascii=False), encoding="utf-8")
    (destination / "sample.jsonl").write_text(json.dumps({"正文": proof + " JSONL资料", "名额": 28}, ensure_ascii=False) + "\n", encoding="utf-8")
    with zipfile.ZipFile(destination / "sample.epub", "w") as book:
        book.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        book.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/book.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        book.writestr("OPS/book.opf", '<?xml version="1.0"?><package version="3.0" xmlns="http://www.idpf.org/2007/opf" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">format-proof</dc:identifier><dc:title>申请资料</dc:title><dc:language>zh</dc:language></metadata><manifest><item id="ch1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="ch1"/></spine></package>')
        book.writestr("OPS/chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>第一章</title></head><body><h1>第一章 申请材料</h1><p>' + proof + ' EPUB资料</p></body></html>')
    if legacy:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            raise RuntimeError("LibreOffice is required to generate genuine DOC/XLS/PPT fixtures")
        for source, target, filter_name in (("docx", "doc", "MS Word 97"), ("xlsx", "xls", "MS Excel 97"), ("pptx", "ppt", "MS PowerPoint 97")):
            with tempfile.TemporaryDirectory(prefix="zc-fixture-office-") as profile:
                subprocess.run([office, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", f"{target}:{filter_name}", "--outdir", str(destination), str(destination / f"sample.{source}")], check=True, capture_output=True, timeout=60)
            if not (destination / f"sample.{target}").is_file():
                raise RuntimeError(f"LibreOffice did not produce a genuine {target} fixture")
    extensions = ["pdf", "docx", "md", "txt", "pptx", "xlsx", "csv", "html", "htm", "json", "jsonl", "doc", "xls", "ppt", "epub"]
    manifest = [{"extension": extension, "filename": f"sample.{extension}", "marker": "formatproof"} for extension in extensions if (destination / f"sample.{extension}").is_file()]
    (destination / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("destination", type=Path)
    parser.add_argument("--legacy", action="store_true")
    args = parser.parse_args()
    print(json.dumps({"files": len(generate(args.destination.resolve(), legacy=args.legacy))}))
