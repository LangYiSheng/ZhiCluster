"""Real HTTP + independent worker smoke test. No AI model or fake data provider.

Run: backend/.venv/Scripts/python scripts/smoke.py [--url URL] [--verify]
Creates four real files and saves IDs for a later restart/persistence verification.
"""
import argparse
import hashlib
import io
import json
import time
from pathlib import Path

import httpx
import pymupdf
from docx import Document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", default="http://127.0.0.1:3000")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    state_path = Path(__file__).resolve().parents[1] / ".data" / "smoke-state.json"
    client = httpx.Client(base_url=args.url, timeout=60)

    def request(method, path, **kwargs):
        response = client.request(method, "/api" + path, **kwargs)
        response.raise_for_status()
        return response

    ready_deadline = time.monotonic() + 60
    while True:
        try:
            if request("GET", "/ready").json()["status"] == "ready":
                break
        except (httpx.HTTPError, ValueError):
            if time.monotonic() >= ready_deadline:
                raise
        time.sleep(0.5)
    if args.verify:
        state = json.loads(state_path.read_text(encoding="utf-8"))
        for item in state["documents"]:
            downloaded = request("GET", f"/documents/{item['id']}/download").content
            assert hashlib.sha256(downloaded).hexdigest() == item["sha256"]
            assert request("GET", f"/documents/{item['id']}").json()["folder_id"] == state["folder_id"]
        print(json.dumps({"status": "PASS", "test": "restart-persistence", "documents": 4}, ensure_ascii=False))
        return
    kb = request("POST", "/knowledge-bases", json={"name": "端到端验收资料", "description": "scripts/smoke.py 生成的真实格式验收文件，可手动删除"}).json()
    kb2 = request("POST", "/knowledge-bases", json={"name": "隔离验证库"}).json()
    folder = request("POST", f"/knowledge-bases/{kb['id']}/folders", json={"name": "研究"}).json()
    child = request("POST", f"/knowledge-bases/{kb['id']}/folders", json={"name": "知识检索", "parent_id": folder["id"]}).json()
    tag = request("POST", f"/knowledge-bases/{kb['id']}/tags", json={"name": "端到端验证", "aliases": ["E2E"]}).json()
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Real knowledge retrieval PDF fixture.")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    word = Document()
    word.add_heading("知识检索", level=1)
    word.add_paragraph("中文检索通过关键词与向量融合实现，引用必须定位到真实原文。")
    buffer = io.BytesIO()
    word.save(buffer)
    files = [("研究.pdf", pdf_bytes), ("研究.docx", buffer.getvalue()), ("研究.md", "# 知识检索\n\n中文检索应隔离知识库，来源引用来自真实文档。".encode()), ("研究.txt", "中文检索的真实文本验收资料。".encode())]
    result = request("POST", f"/knowledge-bases/{kb['id']}/uploads", files=[("files", (name, data)) for name, data in files]).json()
    assert len(result["jobs"]) == 4 and all(item["status"] == "queued" for item in result["items"]), result
    deadline = time.monotonic() + 120
    for job in result["jobs"]:
        while time.monotonic() < deadline:
            status = request("GET", f"/ingestion/jobs/{job['id']}").json()
            if status["status"] in {"completed", "partial_ready", "failed"}:
                assert status["status"] != "failed", status
                break
            time.sleep(0.5)
        else:
            raise AssertionError("Worker processing timed out")
    ids = [item["document"]["id"] for item in result["items"]]
    request("POST", "/documents/batch", json={"kb_id": kb["id"], "document_ids": ids, "action": "move", "folder_id": child["id"]})
    request("POST", "/documents/batch", json={"kb_id": kb["id"], "document_ids": ids, "action": "tags", "tag_ids": [tag["id"]]})
    search = request("POST", "/search", json={"query": "中文检索", "kb_ids": [kb["id"]], "folder_ids": [folder["id"]], "tag_ids": [tag["id"]]}).json()
    assert search["items"] and all(item["kb_id"] == kb["id"] for item in search["items"]), search
    assert not request("POST", "/search", json={"query": "中文检索", "kb_ids": [kb2["id"]]}).json()["items"]
    duplicate = request("POST", f"/knowledge-bases/{kb['id']}/uploads", files=[("files", files[2])]).json()
    assert duplicate["items"][0]["status"] == "duplicate"
    request("DELETE", f"/documents/{ids[0]}")
    assert client.get(f"/api/documents/{ids[0]}/download").status_code == 404
    request("POST", f"/documents/{ids[0]}/restore")
    saved = []
    for doc_id in ids:
        doc = request("GET", f"/documents/{doc_id}").json()
        assert doc["parse_status"] == "ready", doc
        assert request("GET", f"/documents/{doc_id}/sections").json()
        saved.append({"id": doc_id, "sha256": doc["sha256"]})
    events = request("GET", f"/ingestion/jobs/{result['jobs'][0]['id']}/events").text
    assert "event: completed" in events and "id:" in events
    state_path.parent.mkdir(exist_ok=True)
    state_path.write_text(json.dumps({"kb_id": kb["id"], "folder_id": child["id"], "documents": saved}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": "PASS", "test": "real-http-worker-four-formats", "documents": 4, "keyword_hits": len(search["items"]), "mode": search["mode"], "kb_id": kb["id"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
