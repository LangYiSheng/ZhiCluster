"""Regenerate the small real file-format fixtures with backend uv environment."""
from pathlib import Path

import pymupdf
from docx import Document

root = Path(__file__).parent
pdf = pymupdf.open()
pdf.new_page().insert_text((72, 72), "First page: overview of knowledge sources.")
pdf.new_page().insert_text((72, 72), "PositionCheckPDF: actual second-page source anchor.")
pdf.save(root / "anchors.pdf")
doc = Document()
doc.add_heading("Source heading", 0)
doc.add_paragraph("Overview paragraph.")
doc.add_paragraph("PositionCheckDOCX: actual paragraph anchor, no invented pages.")
doc.save(root / "anchors.docx")
